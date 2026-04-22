"""
Print a human-readable dump of one slide for debugging.

Usage:
    python inspect_slide.py deck.pptx 3

Shows every shape's name, bounding box (in inches), and text content. Use
this when a specific slide is reported as broken to see exactly what was
written, without opening PowerPoint.
"""

import sys

from pptx import Presentation


EMU_PER_INCH = 914400


def main():
    if len(sys.argv) < 3:
        print("usage: inspect_slide.py <deck.pptx> <slide_number>")
        sys.exit(2)
    path = sys.argv[1]
    idx = int(sys.argv[2]) - 1

    prs = Presentation(path)
    if idx < 0 or idx >= len(prs.slides):
        print(f"slide {idx + 1} out of range (deck has {len(prs.slides)})")
        sys.exit(1)

    slide = prs.slides[idx]
    sw = prs.slide_width / EMU_PER_INCH
    sh = prs.slide_height / EMU_PER_INCH
    print(f"slide {idx + 1}: {sw:.2f} x {sh:.2f} in, {len(slide.shapes)} shapes")
    print("-" * 72)

    for i, s in enumerate(slide.shapes):
        left = (s.left or 0) / EMU_PER_INCH
        top = (s.top or 0) / EMU_PER_INCH
        width = (s.width or 0) / EMU_PER_INCH
        height = (s.height or 0) / EMU_PER_INCH
        print(
            f"[{i:02d}] {s.shape_type} name={s.name!r} "
            f"pos=({left:.2f},{top:.2f}) size=({width:.2f}x{height:.2f})"
        )
        if s.has_text_frame:
            for j, p in enumerate(s.text_frame.paragraphs):
                if not p.text:
                    continue
                bold = "B" if any(r.font.bold for r in p.runs) else " "
                size = p.font.size.pt if p.font.size else "?"
                print(f"       p{j} [{bold}] size={size} :: {p.text!r}")


if __name__ == "__main__":
    main()

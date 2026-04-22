"""
Reusable helpers for building consultant-style PowerPoint decks with python-pptx.

Import this module at the top of a generator script:

    from pptx_helpers import (
        Palette, Fonts,
        new_presentation, set_slide_bg,
        add_shape, add_rounded_rect, add_oval,
        add_text_box, add_multiline_text,
        text_in_shape, multi_text_in_shape,
        add_slide_title, add_bottom_bar,
        add_right_arrow, add_down_arrow,
    )

All positional arguments (`left`, `top`, `width`, `height`) should be EMU
values produced by `pptx.util.Inches(...)` or `pptx.util.Pt(...)`. Never pass
Python floats directly — that will produce invalid .pptx files.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ---------------------------------------------------------------------------
# Default palette — consultant style (navy + one or two accent colors).
# Override any of these in the generator script by assigning new RGBColor.
# ---------------------------------------------------------------------------
class Palette:
    NAVY = RGBColor(0x1B, 0x2A, 0x4A)
    DARK_NAVY = RGBColor(0x0F, 0x1A, 0x30)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
    MID_GRAY = RGBColor(0x8C, 0x8C, 0x8C)
    DARK_GRAY = RGBColor(0x4A, 0x4A, 0x4A)
    ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xDE)
    ACCENT_TEAL = RGBColor(0x00, 0xA8, 0x8F)
    ACCENT_ORANGE = RGBColor(0xE8, 0x6C, 0x00)
    ACCENT_RED = RGBColor(0xD9, 0x3A, 0x3A)
    LIGHT_BLUE_BG = RGBColor(0xE8, 0xF1, 0xFA)
    LIGHT_TEAL_BG = RGBColor(0xE0, 0xF5, 0xF2)
    BORDER_LIGHT = RGBColor(0xD0, 0xD5, 0xDD)


class Fonts:
    TITLE = "Helvetica Neue"
    BODY = "Helvetica Neue"


# ---------------------------------------------------------------------------
# Presentation setup
# ---------------------------------------------------------------------------
def new_presentation(width_in=13.333, height_in=7.5):
    """Create a blank 16:9 widescreen presentation by default."""
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    return prs


def blank_slide(prs):
    """Add a slide using the blank layout (index 6). Always prefer this over
    placeholder-based layouts to retain full control over positioning."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_slide_bg(slide, color):
    """Fill the entire slide with a solid background color."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


# ---------------------------------------------------------------------------
# Shape primitives
# ---------------------------------------------------------------------------
def _apply_fill_and_border(shape, fill_color, border_color, border_width):
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width or Pt(1)
    else:
        shape.line.fill.background()


def add_shape(slide, left, top, width, height, fill_color=None,
              border_color=None, border_width=None):
    """Rectangle with solid fill and optional border."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _apply_fill_and_border(shape, fill_color, border_color, border_width)
    return shape


def add_rounded_rect(slide, left, top, width, height, fill_color=None,
                     border_color=None, border_width=None, corner=0.05):
    """Rounded rectangle — the canonical building block for consultant cards.
    `corner` controls the roundness (0.0–0.5)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = corner
    _apply_fill_and_border(shape, fill_color, border_color, border_width)
    return shape


def add_oval(slide, left, top, width, height, fill_color=None,
             border_color=None, border_width=None):
    """Oval/circle — useful for numbered step markers."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    _apply_fill_and_border(shape, fill_color, border_color, border_width)
    return shape


def add_right_arrow(slide, left, top, width, height, color=None):
    """Right-pointing arrow shape. Prefer this over `add_connector`."""
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    if color:
        arr.fill.solid()
        arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    return arr


def add_down_arrow(slide, left, top, width, height, color=None):
    """Downward-pointing arrow shape."""
    arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    if color:
        arr.fill.solid()
        arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    return arr


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------
def add_text_box(slide, left, top, width, height, text, font_size=12,
                 font_color=None, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name=None):
    """Single-paragraph text box."""
    font_color = font_color or Palette.DARK_GRAY
    font_name = font_name or Fonts.BODY
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    return box


def add_multiline_text(slide, left, top, width, height, lines,
                       font_size=12, font_color=None, bold=False,
                       alignment=PP_ALIGN.LEFT, font_name=None):
    """Multi-paragraph text box.

    Each item in `lines` is either:
      - a plain string (inherits defaults), or
      - a dict with keys: text, size, color, bold, font, align, indent

    Use this for structured content: label + value, code-like listings,
    or bullet blocks with per-line color/weight variation.
    """
    font_color = font_color or Palette.DARK_GRAY
    font_name = font_name or Fonts.BODY
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, dict):
            p.text = item.get("text", "")
            p.font.size = Pt(item.get("size", font_size))
            p.font.color.rgb = item.get("color", font_color)
            p.font.bold = item.get("bold", bold)
            p.font.name = item.get("font", font_name)
            p.alignment = item.get("align", alignment)
            if "indent" in item:
                p.level = item["indent"]
        else:
            p.text = item
            p.font.size = Pt(font_size)
            p.font.color.rgb = font_color
            p.font.bold = bold
            p.font.name = font_name
            p.alignment = alignment
        p.space_before = Pt(2)
        p.space_after = Pt(2)
    return box


def text_in_shape(shape, text, font_size=11, font_color=None, bold=False,
                  alignment=PP_ALIGN.CENTER, font_name=None):
    """Write a single line of text directly into an existing shape.

    Use this after `add_rounded_rect` / `add_oval` to label a container."""
    font_color = font_color or Palette.DARK_GRAY
    font_name = font_name or Fonts.BODY
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def multi_text_in_shape(shape, lines, font_size=10, font_color=None,
                        alignment=PP_ALIGN.LEFT, font_name=None):
    """Write multiple paragraphs directly into an existing shape with padding."""
    font_color = font_color or Palette.DARK_GRAY
    font_name = font_name or Fonts.BODY
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, dict):
            p.text = item.get("text", "")
            p.font.size = Pt(item.get("size", font_size))
            p.font.color.rgb = item.get("color", font_color)
            p.font.bold = item.get("bold", False)
            p.font.name = item.get("font", font_name)
            p.alignment = item.get("align", alignment)
        else:
            p.text = item
            p.font.size = Pt(font_size)
            p.font.color.rgb = font_color
            p.font.name = font_name
            p.alignment = alignment
        p.space_before = Pt(2)
        p.space_after = Pt(2)


def add_bold_prefix_line(text_frame, label, value, label_color=None,
                         value_color=None, font_size=10, font_name=None,
                         is_first=False):
    """Add a paragraph like `**Label:** value text` — bold prefix + normal tail.

    This is the only reliable way to vary weight within one paragraph in
    python-pptx (you can't use markdown). Pass the parent `text_frame`.
    """
    label_color = label_color or Palette.NAVY
    value_color = value_color or Palette.DARK_GRAY
    font_name = font_name or Fonts.BODY
    p = text_frame.paragraphs[0] if is_first else text_frame.add_paragraph()
    p.space_before = Pt(2)
    p.space_after = Pt(4)
    r1 = p.add_run()
    r1.text = label
    r1.font.size = Pt(font_size)
    r1.font.color.rgb = label_color
    r1.font.bold = True
    r1.font.name = font_name
    r2 = p.add_run()
    r2.text = value
    r2.font.size = Pt(font_size)
    r2.font.color.rgb = value_color
    r2.font.bold = False
    r2.font.name = font_name
    return p


# ---------------------------------------------------------------------------
# Slide chrome (consistent across a deck)
# ---------------------------------------------------------------------------
def add_slide_title(slide, title, subtitle=None, accent_color=None,
                    slide_width_in=13.333):
    """Top bar + title + optional subtitle + short accent rule. Apply on
    every content slide for visual consistency."""
    accent = accent_color or Palette.ACCENT_BLUE
    add_shape(slide, Inches(0), Inches(0), Inches(slide_width_in), Inches(0.06),
              fill_color=accent)
    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(slide_width_in - 2.3),
                 Inches(0.6), title, font_size=26, font_color=Palette.NAVY,
                 bold=True, font_name=Fonts.TITLE)
    if subtitle:
        add_text_box(slide, Inches(0.8), Inches(0.95), Inches(slide_width_in - 2.3),
                     Inches(0.4), subtitle, font_size=13,
                     font_color=Palette.MID_GRAY)
    add_shape(slide, Inches(0.8), Inches(1.35), Inches(1.2), Inches(0.04),
              fill_color=accent)


def add_bottom_bar(slide, page_num, total, deck_title="",
                   slide_width_in=13.333, slide_height_in=7.5,
                   bar_color=None):
    """Navy footer bar with deck title (left) and page counter (right)."""
    bar_color = bar_color or Palette.NAVY
    add_shape(slide, Inches(0), Inches(slide_height_in - 0.4),
              Inches(slide_width_in), Inches(0.4), fill_color=bar_color)
    if deck_title:
        add_text_box(slide, Inches(0.6), Inches(slide_height_in - 0.38),
                     Inches(slide_width_in - 2.5), Inches(0.35),
                     deck_title, font_size=9, font_color=Palette.WHITE)
    add_text_box(slide, Inches(slide_width_in - 1.9),
                 Inches(slide_height_in - 0.38), Inches(1.5), Inches(0.35),
                 f"{page_num} / {total}", font_size=9,
                 font_color=Palette.WHITE, alignment=PP_ALIGN.RIGHT)

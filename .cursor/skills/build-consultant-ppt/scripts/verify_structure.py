"""
Verify the structural integrity of a generated .pptx file.

Run immediately after generation, BEFORE asking the user to open it:

    python verify_structure.py path/to/deck.pptx

Checks performed:
  1. File exists and is non-empty.
  2. .pptx is a valid ZIP archive (OPC container).
  3. Every .xml part parses as well-formed XML (catches the #1 cause of the
     "PowerPoint found a problem with content" error).
  4. python-pptx can open the file.
  5. Expected slide count (optional second arg).
  6. Each slide has at least one shape.

Exits non-zero on any failure so the agent knows to enter the repair loop.
"""

import sys
import os
import zipfile
import xml.etree.ElementTree as ET

from pptx import Presentation


def fail(msg):
    print(f"[FAIL] {msg}")
    sys.exit(1)


def ok(msg):
    print(f"[ OK ] {msg}")


def main():
    if len(sys.argv) < 2:
        print("usage: verify_structure.py <deck.pptx> [expected_slide_count]")
        sys.exit(2)

    path = sys.argv[1]
    expected = int(sys.argv[2]) if len(sys.argv) >= 3 else None

    if not os.path.exists(path):
        fail(f"file not found: {path}")
    if os.path.getsize(path) < 1024:
        fail(f"file too small, likely truncated: {path}")
    ok(f"file exists ({os.path.getsize(path)} bytes)")

    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        fail(f"not a valid .pptx (ZIP) file: {e}")
    ok("valid ZIP container")

    bad = []
    for name in zf.namelist():
        if not name.endswith(".xml") and not name.endswith(".rels"):
            continue
        try:
            ET.fromstring(zf.read(name))
        except ET.ParseError as e:
            bad.append((name, str(e)))
    if bad:
        for name, err in bad:
            print(f"       malformed: {name} -> {err}")
        fail(f"{len(bad)} malformed XML part(s) — this causes 'content problems' on open")
    ok("all XML parts well-formed")

    try:
        prs = Presentation(path)
    except Exception as e:
        fail(f"python-pptx cannot open the file: {e}")
    ok(f"python-pptx opened it, {len(prs.slides)} slide(s)")

    if expected is not None and len(prs.slides) != expected:
        fail(f"expected {expected} slide(s), got {len(prs.slides)}")
    if expected is not None:
        ok(f"slide count matches expected ({expected})")

    empty = [i + 1 for i, s in enumerate(prs.slides) if len(s.shapes) == 0]
    if empty:
        fail(f"empty slide(s) (no shapes): {empty}")
    ok("every slide has at least one shape")

    print("\n[PASS] structural verification succeeded")


if __name__ == "__main__":
    main()
